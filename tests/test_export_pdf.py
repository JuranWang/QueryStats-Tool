from pathlib import Path
import subprocess

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
import pytest

from engine import export_pdf


@pytest.fixture
def sample_docx(tmp_path):
    picture = tmp_path / "chart.png"
    image = Image.new("RGB", (160, 100), "white")
    ImageDraw.Draw(image).rectangle((20, 20, 80, 90), fill="blue")
    image.save(picture)
    document = Document()
    document.add_heading("Conversion test", 0)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph, with a chart below.")
    document.add_picture(str(picture), width=Inches(2))
    source = tmp_path / "报告 sample-Word-0921.docx"
    document.save(source)
    return source


def test_real_pdf_conversion(sample_docx, tmp_path):
    try:
        export_pdf._find_soffice()
    except export_pdf.ExportToolMissingError as exc:
        pytest.skip(str(exc))
    output = tmp_path / "exports"
    output.mkdir()
    # An existing PDF must be replaced by a newly converted report.
    (output / f"{sample_docx.stem}.pdf").write_bytes(b"old report")
    result = Path(export_pdf.convert_docx_to_pdf(str(sample_docx), str(output)))
    assert result.is_absolute()
    assert result == output / f"{sample_docx.stem}.pdf"
    assert result.stat().st_size > 0
    assert result.read_bytes().startswith(b"%PDF")


def test_missing_libreoffice(monkeypatch, sample_docx, tmp_path):
    monkeypatch.setattr(export_pdf.shutil, "which", lambda name: None)
    monkeypatch.setattr(export_pdf, "_is_executable", lambda path: False)
    with pytest.raises(export_pdf.ExportToolMissingError, match="LibreOffice.*soffice.*https://"):
        export_pdf.convert_docx_to_pdf(str(sample_docx), str(tmp_path / "exports"))


@pytest.mark.parametrize("candidate", [
    "/opt/homebrew/bin/soffice", "/usr/local/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
])
def test_mac_fallback(monkeypatch, candidate):
    monkeypatch.setattr(export_pdf.shutil, "which", lambda name: None)
    monkeypatch.setattr(export_pdf, "_is_executable", lambda path: path == candidate)
    assert export_pdf._find_soffice() == candidate


@pytest.mark.parametrize("timeout_error", [False, True])
def test_subprocess_failure(monkeypatch, sample_docx, tmp_path, timeout_error):
    monkeypatch.setattr(export_pdf, "_find_soffice", lambda: "soffice")

    def fail(command, **kwargs):
        assert kwargs["timeout"] == 7
        assert kwargs["check"] is True
        if timeout_error:
            raise subprocess.TimeoutExpired(command, 7, stderr=b"PDF conversion stalled")
        raise subprocess.CalledProcessError(2, command, stderr="PDF conversion failed detail")

    monkeypatch.setattr(export_pdf.subprocess, "run", fail)
    with pytest.raises(RuntimeError, match="PDF conversion") as exc:
        export_pdf.convert_docx_to_pdf(str(sample_docx), str(tmp_path / "exports"), timeout=7)
    assert ("7" if timeout_error else "2") in str(exc.value)
    assert sample_docx.is_file()


def test_success_without_output_does_not_return_stale_pdf(monkeypatch, sample_docx, tmp_path):
    monkeypatch.setattr(export_pdf, "_find_soffice", lambda: "soffice")
    monkeypatch.setattr(export_pdf.subprocess, "run", lambda *a, **kw:
                        subprocess.CompletedProcess(a[0], 0, stdout="", stderr="no PDF produced"))
    stale = tmp_path / f"{sample_docx.stem}.pdf"
    stale.write_bytes(b"old report")
    with pytest.raises(RuntimeError, match="no PDF produced"):
        export_pdf.convert_docx_to_pdf(str(sample_docx), str(tmp_path))
