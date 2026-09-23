from pathlib import Path
import re
import subprocess
from urllib.parse import unquote
import zipfile

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
import pytest

from engine import export_markdown


@pytest.fixture
def sample_docx(tmp_path):
    picture = tmp_path / "chart.png"
    image = Image.new("RGB", (160, 100), "white")
    ImageDraw.Draw(image).rectangle((20, 20, 80, 90), fill="blue")
    image.save(picture)
    document = Document()
    document.add_heading("Conversion test", 0)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph, with a chart below.")
    document.add_picture(str(picture), width=Inches(2))
    source = tmp_path / "报告 sample-Word-0921.docx"
    document.save(source)
    return source


def test_real_markdown_conversion_with_portable_images(sample_docx, tmp_path):
    try:
        export_markdown._find_pandoc()
    except export_markdown.ExportToolMissingError as exc:
        pytest.skip(str(exc))
    output = tmp_path / "exports"
    result = Path(export_markdown.convert_docx_to_markdown(str(sample_docx), str(output)))
    assert result.is_absolute()
    assert result == output / f"{sample_docx.stem}.zip"
    assert result.stat().st_size > 0
    with zipfile.ZipFile(result) as bundle:
        assert bundle.testzip() is None
        assert f"{sample_docx.stem}.md" in bundle.namelist()
        markdown = bundle.read(f"{sample_docx.stem}.md").decode()
        assert "First paragraph" in markdown
        assert str(tmp_path) not in markdown
        images = [name for name in bundle.namelist() if name.endswith(".png")]
        assert images
        assert all(name.startswith(f"{sample_docx.stem}_media/media/") for name in images)
        # Extract elsewhere and resolve every referenced image without the export directory.
        extracted = tmp_path / "another computer"
        bundle.extractall(extracted)
        references = re.findall(r"!\[[^\]]*\]\(<?([^)>]+)>?\)", markdown)
        assert references
        for reference in references:
            relative = Path(unquote(reference))
            assert not relative.is_absolute()
            assert (extracted / relative).is_file()

    # Reusing the same name for a text-only report removes obsolete media.
    document = Document()
    document.add_paragraph("Updated report without images.")
    document.save(sample_docx)
    export_markdown.convert_docx_to_markdown(str(sample_docx), str(output))
    with zipfile.ZipFile(result) as bundle:
        assert not any(name.endswith(".png") for name in bundle.namelist())


def test_missing_pandoc(monkeypatch, sample_docx, tmp_path):
    monkeypatch.setattr(export_markdown.shutil, "which", lambda name: None)
    monkeypatch.setattr(export_markdown, "_is_executable", lambda path: False)
    with pytest.raises(export_markdown.ExportToolMissingError, match="Pandoc.*pandoc.*https://"):
        export_markdown.convert_docx_to_markdown(str(sample_docx), str(tmp_path / "exports"))


@pytest.mark.parametrize("candidate", ["/opt/homebrew/bin/pandoc", "/usr/local/bin/pandoc"])
def test_mac_fallback(monkeypatch, candidate):
    monkeypatch.setattr(export_markdown.shutil, "which", lambda name: None)
    monkeypatch.setattr(export_markdown, "_is_executable", lambda path: path == candidate)
    assert export_markdown._find_pandoc() == candidate


@pytest.mark.parametrize("timeout_error", [False, True])
def test_subprocess_failure(monkeypatch, sample_docx, tmp_path, timeout_error):
    monkeypatch.setattr(export_markdown, "_find_pandoc", lambda: "pandoc")
    output = tmp_path / "exports"

    def fail(command, **kwargs):
        assert kwargs["timeout"] == 7
        assert kwargs["check"] is True
        assert kwargs["cwd"] == output.resolve()
        assert command[command.index("-o") + 1] == f"{sample_docx.stem}.md"
        assert command[-1] == f"--extract-media={sample_docx.stem}_media"
        if timeout_error:
            raise subprocess.TimeoutExpired(command, 7, stderr=b"Pandoc conversion stalled")
        raise subprocess.CalledProcessError(2, command, stderr="Pandoc conversion failed detail")

    monkeypatch.setattr(export_markdown.subprocess, "run", fail)
    with pytest.raises(RuntimeError, match="Pandoc conversion") as exc:
        export_markdown.convert_docx_to_markdown(str(sample_docx), str(output), timeout=7)
    assert ("7" if timeout_error else "2") in str(exc.value)
    assert sample_docx.is_file()


def test_success_without_output_does_not_package_stale_markdown(monkeypatch, sample_docx, tmp_path):
    monkeypatch.setattr(export_markdown, "_find_pandoc", lambda: "pandoc")
    monkeypatch.setattr(export_markdown.subprocess, "run", lambda *a, **kw:
                        subprocess.CompletedProcess(a[0], 0, stdout="", stderr="no Markdown produced"))
    (tmp_path / f"{sample_docx.stem}.md").write_text("old report")
    with pytest.raises(RuntimeError, match="no Markdown produced"):
        export_markdown.convert_docx_to_markdown(str(sample_docx), str(tmp_path))
